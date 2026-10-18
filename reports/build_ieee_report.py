from __future__ import annotations

import csv
import json
import re
from pathlib import Path
from typing import Iterable

import joblib
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "reports" / "IEEE_report_source.md"
ARTIFACTS = ROOT / "artifacts"
DATA = ROOT / "data"
REPORT_DIR = ROOT / "reports"
OUTPUT = ROOT / "I26043443_Dattu_Nayak_Paltya_Report.docx"
FIGURES = REPORT_DIR / "report_figures"
LABELS = ["negative", "neutral", "positive"]
COLORS = {"negative": "#ee7479", "neutral": "#8a79e8", "positive": "#35ae8c"}


def _set_cell_shading(cell, color: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(shading)


def _set_cell_border(cell, **kwargs: dict) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge not in kwargs:
            continue
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        for key in ["val", "sz", "space", "color"]:
            if key in kwargs[edge]:
                element.set(qn(f"w:{key}"), str(kwargs[edge][key]))


def _plain_text(text: str) -> str:
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"\*(.*?)\*", r"\1", text)
    return text.replace("`", "")


def _format_run(run, size: float = 10, bold: bool = False, italic: bool = False) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(25, 25, 25)


def _format_paragraph(paragraph, size: float = 10, align=WD_ALIGN_PARAGRAPH.JUSTIFY) -> None:
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.0
    for run in paragraph.runs:
        _format_run(run, size=size)


def _add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])
    _format_run(run, size=8)


def _set_columns(section, count: int, spacing_twips: int = 360) -> None:
    sect_pr = section._sectPr
    cols = sect_pr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sect_pr.append(cols)
    cols.set(qn("w:num"), str(count))
    cols.set(qn("w:space"), str(spacing_twips))
    cols.set(qn("w:equalWidth"), "1")


def _save_figures() -> None:
    FIGURES.mkdir(parents=True, exist_ok=True)
    distribution = pd.read_csv(ARTIFACTS / "class_distribution.csv")
    fig, axes = plt.subplots(1, 3, figsize=(9, 3.0))
    for ax, split in zip(axes, ["train", "validation", "test"]):
        rows = distribution[distribution["split"] == split].set_index("label")
        values = [int(rows.loc[label, "count"]) for label in LABELS]
        ax.pie(
            values,
            labels=None,
            startangle=90,
            colors=[COLORS[label] for label in LABELS],
            wedgeprops={"linewidth": 1.2, "edgecolor": "white"},
            autopct=lambda pct: f"{pct:.0f}%" if pct >= 7 else "",
            textprops={"fontsize": 8, "color": "white"},
        )
        ax.set_title(f"{split.title()} (n={sum(values):,})", fontsize=9, pad=5)
    fig.legend(
        [plt.Line2D([0], [0], marker="o", color="w", markerfacecolor=COLORS[label], markersize=7) for label in LABELS],
        [label.title() for label in LABELS],
        loc="lower center",
        ncol=3,
        frameon=False,
        fontsize=8,
        bbox_to_anchor=(0.5, -0.02),
    )
    fig.tight_layout(rect=(0, 0.10, 1, 1))
    fig.savefig(FIGURES / "split-distribution.png", dpi=220, bbox_inches="tight")
    plt.close(fig)

    lengths = pd.read_csv(ARTIFACTS / "training_text_length_summary.csv", index_col=0)
    train = pd.read_csv(DATA / "train.csv", usecols=["text"])
    values = train["text"].str.len()
    fig, ax = plt.subplots(figsize=(5.8, 2.8))
    ax.hist(values, bins=35, color="#8a79e8", edgecolor="white", linewidth=0.7)
    ax.axvline(values.median(), color="#172033", linestyle="--", linewidth=1, label=f"Median {values.median():.0f}")
    ax.set_xlabel("Characters per tweet")
    ax.set_ylabel("Training examples")
    ax.set_title("Tweet length distribution", fontsize=10)
    ax.grid(axis="y", alpha=0.2)
    ax.legend(frameon=False, fontsize=8)
    fig.tight_layout()
    fig.savefig(FIGURES / "text-length.png", dpi=220, bbox_inches="tight")
    plt.close(fig)

    metrics = pd.read_csv(ARTIFACTS / "metrics.csv")
    cv = pd.read_csv(ARTIFACTS / "cross_validation_results.csv")
    model_order = ["majority_baseline", "multinomial_nb", "logistic_regression", "linear_svm"]
    labels = ["Majority", "Naive Bayes", "Logistic Reg.", "Linear SVM"]
    x = np.arange(len(model_order))
    width = 0.19
    fig, ax = plt.subplots(figsize=(6.0, 3.2))
    for j, (metric, name, color) in enumerate(
        [
            ("accuracy", "Validation accuracy", "#9b8af0"),
            ("test_accuracy", "Test accuracy", "#53b99a"),
            ("macro_f1", "Validation macro-F1", "#e9aa5b"),
            ("test_macro_f1", "Test macro-F1", "#5b91d1"),
        ]
    ):
        vals = [float(metrics.set_index("model").loc[m, metric]) for m in model_order]
        ax.bar(x + (j - 1.5) * width, vals, width, label=name, color=color)
    ax.set_xticks(x, labels, rotation=12, ha="right")
    ax.set_ylim(0, 0.82)
    ax.set_ylabel("Score")
    ax.set_title("Model performance across official splits", fontsize=10)
    ax.grid(axis="y", alpha=0.22)
    ax.legend(frameon=False, fontsize=7, ncol=2)
    fig.tight_layout()
    fig.savefig(FIGURES / "model-comparison.png", dpi=220, bbox_inches="tight")
    plt.close(fig)

    predictions = pd.read_csv(ARTIFACTS / "test_predictions.csv")
    matrix = pd.crosstab(
        pd.Categorical(predictions["true_label"], LABELS),
        pd.Categorical(predictions["predicted_label"], LABELS),
        dropna=False,
    ).reindex(index=LABELS, columns=LABELS, fill_value=0)
    fig, ax = plt.subplots(figsize=(4.4, 3.3))
    image = ax.imshow(matrix.to_numpy(), cmap="Purples")
    ax.set_xticks(range(3), [label.title() for label in LABELS])
    ax.set_yticks(range(3), [label.title() for label in LABELS])
    ax.set_xlabel("Predicted label")
    ax.set_ylabel("Actual label")
    ax.set_title("Linear SVM test confusion matrix", fontsize=9)
    cutoff = matrix.to_numpy().max() * 0.58
    for i in range(3):
        for j in range(3):
            value = int(matrix.iloc[i, j])
            ax.text(j, i, f"{value:,}", ha="center", va="center", fontsize=8,
                    color="white" if value > cutoff else "#302b45")
    fig.colorbar(image, ax=ax, fraction=0.045, pad=0.04)
    fig.tight_layout()
    fig.savefig(FIGURES / "confusion-matrix.png", dpi=220, bbox_inches="tight")
    plt.close(fig)

    model = joblib.load(ARTIFACTS / "sentiment_model.joblib")
    tfidf = model.named_steps["tfidf"]
    classifier = model.named_steps["classifier"]
    terms = tfidf.get_feature_names_out()
    rows: list[dict[str, object]] = []
    for class_index, label_id in enumerate(classifier.classes_):
        weights = classifier.coef_[class_index]
        indices = np.argsort(np.abs(weights))[-8:][::-1]
        for index in indices:
            rows.append({"class": LABELS[int(label_id)], "feature": terms[index], "coefficient": float(weights[index])})
    pd.DataFrame(rows).to_csv(FIGURES / "linear_svm_coefficients.csv", index=False)
    terms_frame = pd.DataFrame(rows)
    fig, axes = plt.subplots(1, 3, figsize=(9, 3.6))
    for ax, label in zip(axes, LABELS):
        part = terms_frame[terms_frame["class"] == label].sort_values("coefficient")
        ax.barh(part["feature"], part["coefficient"],
                color=COLORS[label], height=0.72)
        ax.set_title(label.title(), fontsize=9)
        ax.tick_params(axis="y", labelsize=6)
        ax.tick_params(axis="x", labelsize=7)
        ax.axvline(0, color="#50576a", linewidth=0.7)
        ax.grid(axis="x", alpha=0.18)
    fig.suptitle("Largest-magnitude Linear SVM coefficients", fontsize=10)
    fig.tight_layout()
    fig.savefig(FIGURES / "top-features.png", dpi=220, bbox_inches="tight")
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(6.0, 3.1))
    colors = {"multinomial_nb": "#83b6ec", "logistic_regression": "#9a88ec", "linear_svm": "#55b9a0"}
    offsets = {"multinomial_nb": -0.22, "logistic_regression": 0.0, "linear_svm": 0.22}
    for model_name in colors:
        part = cv[cv["model"] == model_name]
        row = part.loc[part["mean_macro_f1"].idxmax()]
        mean = float(row["mean_macro_f1"])
        scores = [float(row[f"fold_{fold}_macro_f1"]) for fold in range(3)]
        y = mean
        x = ["Naive Bayes", "Logistic Reg.", "Linear SVM"][["multinomial_nb", "logistic_regression", "linear_svm"].index(model_name)]
        jitter = offsets[model_name]
        ax.errorbar(x, mean, yerr=float(row["std_macro_f1"]), fmt="o",
                    color=colors[model_name], capsize=4, markersize=7,
                    label=f"{model_name.replace('_', ' ')} mean")
        ax.scatter([x] * len(scores), scores, color=colors[model_name], s=18, alpha=0.75, zorder=3)
    ax.set_ylim(0.58, 0.69)
    ax.set_ylabel("Macro-F1")
    ax.set_title("Best training CV configurations (three folds)", fontsize=10)
    ax.grid(axis="y", alpha=0.2)
    fig.tight_layout()
    fig.savefig(FIGURES / "cv-folds.png", dpi=220, bbox_inches="tight")
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(6.0, 2.6))
    ax.axis("off")
    boxes = [
        ("TweetEval", "official train / validation / test"),
        ("Inspect", "labels · duplicates · text quality"),
        ("TF-IDF", "word + character n-grams"),
        ("Compare", "3-fold CV · four model families"),
        ("Evaluate", "validation selection · held-out test"),
        ("Prototype", "saved model · local dashboard"),
    ]
    x_centers = np.linspace(0.09, 0.91, len(boxes))
    for index, ((title, sub), x) in enumerate(zip(boxes, x_centers)):
        ax.text(x, 0.58, f"{title}\n{sub}", ha="center", va="center", fontsize=8,
                bbox={"boxstyle": "round,pad=0.65", "facecolor": "#f4f1ff" if index % 2 == 0 else "#eff9f5",
                      "edgecolor": "#c9c2ed" if index % 2 == 0 else "#b8e3d4"})
        if index < len(boxes) - 1:
            ax.annotate("", xy=(x_centers[index + 1] - 0.055, 0.58),
                        xytext=(x + 0.055, 0.58),
                        arrowprops={"arrowstyle": "->", "color": "#8a8fa0", "lw": 1.0})
    ax.set_title("Reproducible sentiment-classification workflow", fontsize=10, pad=8)
    fig.tight_layout()
    fig.savefig(FIGURES / "system-architecture.png", dpi=220, bbox_inches="tight")
    plt.close(fig)


def _add_markdown_table(document, lines: list[str]) -> None:
    rows = []
    for line in lines:
        if re.match(r"^\|\s*:?-{2,}", line):
            continue
        rows.append([cell.strip() for cell in line.strip().strip("|").split("|")])
    if not rows:
        return
    col_count = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=col_count)
    table.autofit = True
    for row_index, row in enumerate(rows):
        for col_index in range(col_count):
            cell = table.cell(row_index, col_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = _plain_text(row[col_index]) if col_index < len(row) else ""
            if row_index == 0:
                _set_cell_shading(cell, "EDE9FB")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(1)
                paragraph.paragraph_format.space_before = Pt(1)
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in paragraph.runs:
                    _format_run(run, size=7.3, bold=row_index == 0)
            border = {"val": "single", "sz": "3", "color": "D9DDE6"}
            _set_cell_border(cell, top=border, bottom=border, left=border, right=border)
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def _add_figure(document, filename: str, caption: str) -> None:
    path = FIGURES / filename
    if not path.is_file():
        raise FileNotFoundError(f"Report figure is missing: {path}")
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.add_run().add_picture(str(path), width=Inches(3.25))
    caption_paragraph = document.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.space_after = Pt(5)
    caption_run = caption_paragraph.add_run(caption)
    _format_run(caption_run, size=8)


def _add_paragraph(document, text: str, *, size: float = 10, bold_prefix: str | None = None) -> None:
    text = _plain_text(text)
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.keep_together = False
    paragraph.paragraph_format.widow_control = True
    paragraph.paragraph_format.space_after = Pt(3 if size >= 9 else 1)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        _format_run(first, size=size, bold=True)
        run = paragraph.add_run(text[len(bold_prefix):])
        _format_run(run, size=size)
    else:
        run = paragraph.add_run(text)
        _format_run(run, size=size)


def _write_docx() -> int:
    _save_figures()
    content = SOURCE.read_text(encoding="utf-8").splitlines()
    document = Document()
    for section in document.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(0.625)
        section.right_margin = Inches(0.625)
        section.header_distance = Inches(0.3)
        section.footer_distance = Inches(0.35)
        _set_columns(section, 1)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.0
    for style_name in ["Heading 1", "Heading 2"]:
        style = document.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.color.rgb = RGBColor(25, 25, 25)

    title = content[0].removeprefix("# ").strip()
    title_paragraph = document.add_paragraph()
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_paragraph.paragraph_format.space_before = Pt(0)
    title_paragraph.paragraph_format.space_after = Pt(9)
    title_paragraph.paragraph_format.keep_with_next = True
    run = title_paragraph.add_run(title)
    _format_run(run, size=24, bold=False)

    author = document.add_paragraph()
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author.paragraph_format.space_after = Pt(2)
    author.paragraph_format.keep_with_next = True
    _format_run(author.add_run("Dattu Nayak Paltya"), size=9)
    affiliation = document.add_paragraph()
    affiliation.alignment = WD_ALIGN_PARAGRAPH.CENTER
    affiliation.paragraph_format.space_after = Pt(2)
    affiliation.paragraph_format.keep_with_next = True
    _format_run(
        affiliation.add_run("BCSI Programme, INTI International University, Malaysia"),
        size=9,
    )
    course = document.add_paragraph()
    course.alignment = WD_ALIGN_PARAGRAPH.CENTER
    course.paragraph_format.space_after = Pt(7)
    _format_run(course.add_run("PRG4206 Machine Learning · Student ID I26043443"), size=9)

    two_column_section = document.add_section(WD_SECTION_START.CONTINUOUS)
    two_column_section.page_width = Inches(8.27)
    two_column_section.page_height = Inches(11.69)
    two_column_section.top_margin = Inches(0.75)
    two_column_section.bottom_margin = Inches(1.0)
    two_column_section.left_margin = Inches(0.625)
    two_column_section.right_margin = Inches(0.625)
    two_column_section.header_distance = Inches(0.3)
    two_column_section.footer_distance = Inches(0.35)
    _set_columns(two_column_section, 2, 360)
    _add_page_number(two_column_section.footer.paragraphs[0])

    words = 0
    index = content.index("## Abstract")
    while index < len(content):
        line = content[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("## "):
            heading = line[3:].strip()
            if heading == "Abstract":
                index += 1
                while index < len(content) and not content[index].strip():
                    index += 1
                abstract = content[index].strip()
                _add_paragraph(document, f"Abstract—{abstract}", size=9, bold_prefix="Abstract—")
                words += len(abstract.split())
                index += 1
                continue
            if heading == "ACKNOWLEDGMENT AND AI-USE DISCLOSURE":
                para = document.add_paragraph()
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                para.paragraph_format.keep_with_next = True
                para.paragraph_format.space_before = Pt(5)
                para.paragraph_format.space_after = Pt(4)
                _format_run(para.add_run(heading), size=10, bold=True)
                index += 1
                continue
            para = document.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.keep_with_next = True
            para.paragraph_format.space_before = Pt(5)
            para.paragraph_format.space_after = Pt(4)
            _format_run(para.add_run(heading), size=10, bold=True)
            index += 1
            continue
        if line.startswith("### "):
            para = document.add_paragraph()
            para.paragraph_format.keep_with_next = True
            para.paragraph_format.space_before = Pt(3)
            para.paragraph_format.space_after = Pt(2)
            _format_run(para.add_run(line[4:].strip()), size=10, bold=True, italic=True)
            index += 1
            continue
        if line.startswith("**Keywords—**"):
            text = line.replace("**Keywords—**", "Keywords—", 1)
            _add_paragraph(document, text, size=9, bold_prefix="Keywords—")
            words += len(text.split())
            index += 1
            continue
        if line.startswith("**TABLE "):
            caption = line.replace("**", "").strip()
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.space_before = Pt(3)
            paragraph.paragraph_format.space_after = Pt(3)
            _format_run(paragraph.add_run(caption), size=8, bold=True)
            index += 1
            while index < len(content) and not content[index].strip():
                index += 1
            table_lines = []
            while index < len(content) and content[index].strip().startswith("|"):
                table_lines.append(content[index].strip())
                index += 1
            _add_markdown_table(document, table_lines)
            continue
        if line.startswith("**[FIGURE:"):
            match = re.match(r"\*\*\[FIGURE:([^|]+)\|(.+)\]\*\*", line)
            if not match:
                raise ValueError(f"Malformed figure directive: {line}")
            _add_figure(document, match.group(1).strip(), match.group(2).strip())
            index += 1
            continue
        if line.startswith("[") and re.match(r"^\[\d+\]", line):
            _add_paragraph(document, line, size=8)
            words += len(line.split())
            index += 1
            continue
        if line.startswith("## ") or line.startswith("### "):
            continue
        paragraph_lines = [line]
        index += 1
        while index < len(content):
            following = content[index].strip()
            if (
                not following
                or following.startswith("#")
                or following.startswith("**[FIGURE:")
                or following.startswith("**TABLE ")
                or following.startswith("|")
                or re.match(r"^\[\d+\]", following)
            ):
                break
            paragraph_lines.append(following)
            index += 1
        text = " ".join(paragraph_lines)
        _add_paragraph(document, text, size=9 if text.startswith("Keywords—") else 10)
        words += len(text.split())

    document.core_properties.title = title
    document.core_properties.author = "Dattu Nayak Paltya"
    document.core_properties.subject = "PRG4206 Machine Learning individual assignment"
    document.core_properties.keywords = "TweetEval, sentiment analysis, machine learning, IEEE"
    document.save(OUTPUT)
    return words


def main() -> None:
    word_count = _write_docx()
    print(f"Created {OUTPUT}")
    print(f"Approximate report word count: {word_count:,}")
    print(f"Figures: {FIGURES}")


if __name__ == "__main__":
    main()
